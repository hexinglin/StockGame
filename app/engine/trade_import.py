"""
模块名称: engine/trade_import.py
说明:    成交记录导出解析（页面导入通道：粘贴文本 / CSV / Excel 工作簿）

将券商客户端 / QMT 客户端导出的成交明细解析为标准成交记录（与 agent 上报
结构一致，直接复用 trade_analysis 分析）。

支持:
- 输入通道: 文本（表格复制粘贴 / CSV）与 Excel 工作簿（.xlsx，券商"历史成交
  查询"导出，前几行为营业部/账号等元信息，表头在其后）；
- 分隔符自动识别：Tab（客户端右键复制）优先，其次逗号（CSV 导出）；
- 表头自动识别：在前 30 行内找含"成交/价格/数量/代码"等关键词 ≥2 个的表头
  行（自动跳过券商导出的元信息头），按列名模糊匹配（高优先列名先行，低优先
  兜底）：日期/时间/代码/名称/方向/价格/数量/金额/成交编号/委托编号/交易市场；
- 日期+时间：支持合一列（'YYYY-MM-DD HH:MM:SS'，含 '/' 分隔与毫秒）与拆分
  两列（券商导出 '20260911' + '14:52:44.03' 毫秒成交时间）两种形态；
- 方向：证券买入/融资买入/买入/买/48/23 → buy；融资卖出还款/卖出/49/24 →
  sell；交易类别含"上海/深圳"前缀时据此推断市场；
- 代码：'588000' 无后缀时按交易市场列或交易类别前缀补 .SH/.SZ；
- 多日：一次导出可跨多个交易日，记录按自带日期分组（group_by_date），
  无日期列的行回退所选日期；
- 容错：跳过空行/合计行/脏行，逐行原因返回 skipped 供页面展示。
"""
import re

# 高优先列名（更精确，先行匹配）
_HIGH = (
    ("time", ("成交时间",)),
    ("date", ("成交日期", "交易日期")),
    ("price", ("成交价格", "成交均价")),
    ("volume", ("成交数量", "成交量")),
    ("amount", ("成交金额",)),
    ("code", ("证券代码", "股票代码")),
    ("trade_id", ("成交编号",)),
)
# 低优先列名（兜底；date 在 time 之前，避免"日期"列被 time 抢占）
_LOOSE = (
    ("date", ("日期",)),
    ("time", ("时间",)),
    ("price", ("价格",)),
    ("volume", ("数量",)),
    ("amount", ("金额", "成交额")),
    ("code", ("代码", "证券编码")),
    ("name", ("名称",)),
    ("direction", ("买卖", "操作", "方向", "类别")),
    ("trade_id", ("成交号",)),
    ("order_id", ("委托编号", "合同编号", "委托号", "委托序号")),
    ("market", ("市场", "交易所")),
)
# 表头行判定关键词（行内含 >=2 个即视为表头候选）
_HEADER_HINTS = ("成交", "价格", "数量", "代码", "时间", "日期", "金额",
                 "名称", "买卖", "方向")

_MARKET_MAP = (
    ("上海", ".SH"), ("沪", ".SH"), ("SH", ".SH"), ("1", ".SH"),
    ("深圳", ".SZ"), ("深", ".SZ"), ("SZ", ".SZ"), ("0", ".SZ"),
)


def _split_line(line):
    """按分隔符拆分一行：Tab 优先，其次逗号，最后 2+ 连续空格"""
    if "\t" in line:
        return [c.strip() for c in line.split("\t")]
    if "," in line:
        return [c.strip().strip('"') for c in line.split(",")]
    if re.search(r"\s{2,}", line):
        return [c.strip() for c in re.split(r"\s{2,}", line)]
    return [line.strip()]


def _find_header_row(cell_rows):
    """在细胞行序列中找表头行（跳过券商导出的元信息头），返回行下标或 None"""
    best = None
    for i, cells in enumerate(cell_rows[:30]):
        if len(cells) < 3:
            continue
        hit = sum(1 for c in cells
                  if any(h in str(c) for h in _HEADER_HINTS))
        if hit >= 2 and (best is None or hit > best[1]):
            best = (i, hit)
    return best[0] if best is not None else None


def _match_header(cells):
    """表头单元格 → 字段名映射（高优先先行，低优先兜底，同字段先到先得）"""
    mapping = {}
    assigned = set()

    def _try(table):
        for idx, cell in enumerate(cells):
            if idx in mapping:
                continue
            c = str(cell)
            for field, kws in table:
                if field in assigned:
                    continue
                if any(kw in c for kw in kws):
                    mapping[idx] = field
                    assigned.add(field)
                    break

    _try(_HIGH)
    _try(_LOOSE)
    return mapping


def _norm_date(v):
    """日期 → 'YYYY-MM-DD'；支持 20260911 / 2026-09-11 / 2026/9/11 / datetime；失败空串"""
    s = str(v or "").strip()
    if not s:
        return ""
    m = re.match(r"^(\d{4})[-/.]?(\d{1,2})[-/.]?(\d{1,2})", s.replace("/", "-"))
    if m:
        return "%s-%02d-%02d" % (m.group(1), int(m.group(2)), int(m.group(3)))
    return ""


def _norm_time(v):
    """时间部分 → 'HH:MM:SS'；支持毫秒 '14:52:44.03' / '9:31:15' / '093115' / '09:31'"""
    s = str(v or "").strip()
    if not s:
        return ""
    m = re.match(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?(?:\.\d+)?$", s)
    if m:
        return "%02d:%s:%s" % (int(m.group(1)), m.group(2), m.group(3) or "00")
    if s.isdigit() and 1 <= len(s) <= 6:
        t6 = s.rjust(6, "0")
        return "%s:%s:%s" % (t6[:2], t6[2:4], t6[4:6])
    return ""


def _norm_dt(v, date):
    """时间列 → 'YYYY-MM-DD HH:MM:SS'；含完整日期直取（容忍毫秒），
    仅时间部分时结合 date；失败返回空串"""
    t = str(v or "").strip()
    if not t:
        return ""
    t = t.replace("/", "-").replace("T", " ").strip()
    # '日期 时间' 两段形态（含毫秒尾缀）
    m = re.match(r"^(\S+)[ ]+(\S+)$", t)
    if m:
        d = _norm_date(m.group(1))
        tm = _norm_time(m.group(2))
        if d and tm:
            return "%s %s" % (d, tm)
        if d and not tm and ":" not in m.group(2):
            return "%s 00:00:00" % d       # 仅日期（无时间部分）
    # 紧凑 'YYYYMMDDHHMMSS'
    if t.isdigit() and len(t) == 14:
        return "%s-%s-%s %s:%s:%s" % (t[:4], t[4:6], t[6:8],
                                      t[8:10], t[10:12], t[12:14])
    # 仅时间 'H:MM:SS' / 'HH:MM' / '93115' / '093115'（含毫秒）
    tm = _norm_time(t)
    if tm and date:
        return "%s %s" % (date, tm)
    return ""


def _norm_direction(v):
    """方向归一：买入类→buy，卖出类→sell，其余 unknown"""
    s = str(v or "").strip()
    if not s:
        return "unknown"
    if s in ("48", "23"):
        return "buy"
    if s in ("49", "24"):
        return "sell"
    if "买" in s:
        return "buy"
    if "卖" in s:
        return "sell"
    low = s.lower()
    if low in ("buy", "b"):
        return "buy"
    if low in ("sell", "s"):
        return "sell"
    return "unknown"


def _norm_code(code, market):
    """代码归一：无后缀且有市场信息（市场列/交易类别）时补 .SH/.SZ"""
    c = str(code or "").strip().upper()
    if not c:
        return ""
    if "." in c:
        return c
    m = str(market or "").strip()
    if m:
        for kw, suffix in _MARKET_MAP:
            if kw in m:
                return c + suffix
    return c


def _parse_rows(header_cells, data_rows, fallback_date):
    """通用行解析：表头细胞 + 数据细胞行序列 → (records, skipped)

    data_rows: [(行号, cells)]，行号用于 skipped 展示（文本 1 起行号 /
    Excel 实际行号）；无日期列的行回退 fallback_date。
    """
    mapping = _match_header(header_cells)
    if "code" not in mapping.values() or "price" not in mapping.values() \
            or "volume" not in mapping.values():
        raise ValueError(
            "表头缺少必需列（证券代码/成交价格/成交数量），识别到的列: %s"
            % list(mapping.values()))

    records = []
    skipped = []

    def _val(field, cells):
        for idx, f in mapping.items():
            if f == field and idx < len(cells):
                return cells[idx]
        return ""

    for lineno, cells in data_rows:
        joined = " ".join(str(c) for c in cells)
        if not joined.strip():
            continue
        # 合计/汇总行
        if any(k in joined for k in ("合计", "总计", "汇总")) \
                and not _norm_date(_val("date", cells)):
            skipped.append({"line": lineno, "reason": "合计/汇总行",
                            "text": joined[:80]})
            continue

        # 日期：行内日期列优先（多日导出按各自日期分组落库），否则回退所选日期
        row_date = _norm_date(_val("date", cells)) or fallback_date
        rec_time = _norm_dt(_val("time", cells), row_date)
        if not rec_time and row_date:
            rec_time = "%s 00:00:00" % row_date
        if not rec_time:
            skipped.append({"line": lineno,
                            "reason": "日期缺失（无日期列且未选导入日期）",
                            "text": joined[:80]})
            continue
        code = _norm_code(_val("code", cells),
                          _val("market", cells) or _val("direction", cells))
        try:
            volume = int(float(str(_val("volume", cells)).replace(",", "") or 0))
        except ValueError:
            volume = 0
        try:
            price = float(str(_val("price", cells)).replace(",", "") or 0)
        except ValueError:
            price = 0.0
        try:
            amount = float(str(_val("amount", cells)).replace(",", "") or 0)
        except ValueError:
            amount = 0.0
        if not code or volume <= 0:
            skipped.append({"line": lineno, "reason": "代码或数量缺失",
                            "text": joined[:80]})
            continue
        if amount <= 0:
            amount = round(price * volume, 2)
        records.append({
            "time": rec_time,
            "code": code,
            "name": str(_val("name", cells) or "")[:50],
            "direction": _norm_direction(_val("direction", cells)),
            "price": price,
            "volume": volume,
            "amount": round(amount, 2),
            "trade_id": str(_val("trade_id", cells) or "")[:50],
            "order_id": str(_val("order_id", cells) or "")[:50],
        })

    return records, skipped


def parse_export_text(text, date=""):
    """解析导出文本 → (records, meta)

    records: 标准成交记录列表（字段与 agent 上报一致，time 自带完整日期）
    meta: {parsed, skipped, header_line, total_lines}
          skipped: [{'line': 行号(1 起), 'reason': 原因, 'text': 原文摘要}]
    """
    records = []
    skipped = []
    lines = [ln for ln in str(text or "").replace("\r\n", "\n").split("\n")]
    if not any(ln.strip() for ln in lines):
        return [], {"parsed": 0, "skipped": skipped,
                    "header_line": 0, "total_lines": 0}

    header_idx = _find_header_row([_split_line(ln) for ln in lines])
    if header_idx is None:
        raise ValueError(
            "未找到表头行（需包含 成交/价格/数量/代码 等列名；"
            "请从客户端表格右键复制后粘贴，或导出 CSV / Excel 选择文件）")
    data_rows = [(i + 1, _split_line(lines[i]))
                 for i in range(header_idx + 1, len(lines))]
    records, skipped = _parse_rows(_split_line(lines[header_idx]),
                                   data_rows, _norm_date(date))

    return records, {"parsed": len(records), "skipped": skipped,
                     "header_line": header_idx + 1, "total_lines": len(lines)}


def _cell_str(v):
    """Excel 单元格 → 字符串（None→''；整数浮点去掉 .0；其余 str）"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_export_workbook(stream):
    """解析 Excel 工作簿（.xlsx 文件流或路径）→ (records, meta)

    券商"历史成交查询"导出：前几行为营业部/账号等元信息，表头行之后为
    数据行；元信息行不含表头关键词，自动跳过。meta 的 line 为实际行号。
    """
    try:
        import openpyxl
    except ImportError:
        raise ValueError("服务端未安装 openpyxl，无法解析 Excel 文件"
                         "（pip install openpyxl）")
    wb = openpyxl.load_workbook(stream, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        cell_rows = [[_cell_str(v) for v in row]
                     for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    header_idx = _find_header_row(cell_rows)
    if header_idx is None:
        raise ValueError(
            "Excel 中未找到表头行（需包含 成交/价格/数量/代码 等列名）")
    data_rows = [(i + 1, cells) for i, cells in
                 enumerate(cell_rows[header_idx + 1:], header_idx + 1)]
    records, skipped = _parse_rows(cell_rows[header_idx], data_rows, "")

    return records, {"parsed": len(records), "skipped": skipped,
                     "header_line": header_idx + 1, "total_lines": len(cell_rows)}


def group_by_date(records):
    """记录按交易日分组 → {'YYYY-MM-DD': [records]}（组内按成交时间升序）"""
    days = {}
    for r in records:
        days.setdefault(r["time"][:10], []).append(r)
    return {d: sorted(rs, key=lambda r: r["time"]) for d, rs in days.items()}
